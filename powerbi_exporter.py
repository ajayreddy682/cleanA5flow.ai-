"""Power BI-ready model preparation and export package.

What this module creates:

* A suggested star schema (fact + dimension tables) when the data supports it.
* CSV files (UTF-8-BOM) for every model table.
* An Excel workbook containing the same tables plus model documentation.
* Relationship, data-dictionary, KPI, and DAX recommendations.
* A written import guide for Power BI Desktop.

What it deliberately does **not** create: a `.pbix` report. There is no
supported free API that writes arbitrary PBIX files. The package is a dataset
and model blueprint that the user opens in Power BI Desktop.
"""

from __future__ import annotations

import io
import json
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from modules.data_cleaner import CleaningResult
from modules.data_profiler import (
    ColumnKind,
    ColumnProfile,
    format_number,
    profile_dataframe,
)
from modules.report_generator import dataframe_to_csv
from utils.security import is_formula_like
from utils.validators import safe_sheet_name

APP_VERSION = "1.0.0"
NAVY = "0B1F3A"
ROYAL = "1D4ED8"
CYAN = "0891B2"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Calibri", color="1E293B", size=10)


@dataclass
class ModelTable:
    name: str
    dataframe: pd.DataFrame
    role: str  # Fact | Dimension | Flat table
    source_table: str
    description: str = ""
    profiles: list[ColumnProfile] = field(default_factory=list)


@dataclass
class Relationship:
    from_table: str
    from_column: str
    to_table: str
    to_column: str
    cardinality: str
    cross_filter: str
    note: str = ""

    def to_dict(self) -> dict[str, str]:
        return {
            "From table": self.from_table,
            "From column": self.from_column,
            "To table": self.to_table,
            "To column": self.to_column,
            "Cardinality": self.cardinality,
            "Cross-filter direction": self.cross_filter,
            "Note": self.note,
        }


@dataclass
class DaxMeasure:
    name: str
    table: str
    expression: str
    format_string: str
    description: str


@dataclass
class VisualRecommendation:
    visual: str
    title: str
    fields: str
    why: str

    def to_dict(self) -> dict[str, str]:
        return {"Visual": self.visual, "Title": self.title, "Fields": self.fields, "Why": self.why}


@dataclass
class PowerBIModel:
    tables: list[ModelTable]
    relationships: list[Relationship]
    measures: list[DaxMeasure]
    data_dictionary: pd.DataFrame
    kpi_recommendations: pd.DataFrame
    visual_recommendations: list[VisualRecommendation]
    notes: list[str]
    source_tables: list[str]
    fact_tables: list[str] = field(default_factory=list)
    dimension_tables: list[str] = field(default_factory=list)

    @property
    def model_frame(self) -> pd.DataFrame:
        rows = [
            {
                "Table": table.name,
                "Role": table.role,
                "Rows": len(table.dataframe),
                "Columns": table.dataframe.shape[1],
                "Source": table.source_table,
                "Purpose": table.description,
            }
            for table in self.tables
        ]
        return pd.DataFrame(rows)

    @property
    def relationships_frame(self) -> pd.DataFrame:
        if not self.relationships:
            return pd.DataFrame(columns=["From table", "From column", "To table", "To column", "Cardinality", "Cross-filter direction", "Note"])
        return pd.DataFrame([relationship.to_dict() for relationship in self.relationships])

    @property
    def measures_frame(self) -> pd.DataFrame:
        if not self.measures:
            return pd.DataFrame(columns=["Measure", "Table", "DAX expression", "Format string", "Description"])
        return pd.DataFrame(
            [
                {
                    "Measure": measure.name,
                    "Table": measure.table,
                    "DAX expression": measure.expression,
                    "Format string": measure.format_string,
                    "Description": measure.description,
                }
                for measure in self.measures
            ]
        )

    @property
    def visuals_frame(self) -> pd.DataFrame:
        return pd.DataFrame([visual.to_dict() for visual in self.visual_recommendations])

    def dax_text(self) -> str:
        lines = [
            "// CleanFlow AI — suggested DAX measures",
            "// Paste each measure into Power BI Desktop: Modeling > New measure.",
            "// Review table and column names after import; rename if your import differs.",
            "// Dates are stored as date types. Mark Dim_Date as a date table before time intelligence.",
            "",
        ]
        for measure in self.measures:
            lines.append(f"// {measure.description}")
            lines.append(f"{measure.name} = {measure.expression}")
            lines.append(f"// Suggested format string: {measure.format_string}")
            lines.append("")
        return "\n".join(lines)


@dataclass
class PowerBIPackage:
    model: PowerBIModel
    zip_bytes: bytes
    workbook_bytes: bytes
    dax_text: str
    import_guide: str
    dictionary_csv: bytes
    relationships_csv: bytes


# ---------------------------------------------------------------------------
# Model construction
# ---------------------------------------------------------------------------
def prepare_model(
    results: dict[str, CleaningResult],
    *,
    max_dimensions_per_table: int = 5,
) -> PowerBIModel:
    """Build a suggested Power BI model from one or more cleaned tables."""

    if not results:
        raise ValueError("No cleaned tables are available for a Power BI model.")

    tables: list[ModelTable] = []
    relationships: list[Relationship] = []
    notes: list[str] = [
        "This package contains Power BI-ready tables, a suggested model, and DAX suggestions.",
        "It is not a .pbix report file. Power BI Desktop is required to build and save the report.",
        "Relationship and measure suggestions are based on column names, types, and cardinality. Review them in Model view.",
    ]
    source_tables = list(results.keys())
    multi_source = len(results) > 1
    used_table_names: set[str] = set()

    for source_name, result in results.items():
        cleaned = result.cleaned
        if cleaned.empty or cleaned.shape[1] == 0:
            notes.append(f"Table '{source_name}' is empty after cleaning and was skipped.")
            continue
        built_tables, built_relationships = _build_source_model(
            source_name,
            result,
            multi_source=multi_source,
            max_dimensions=max_dimensions_per_table,
            used_names=used_table_names,
        )
        tables.extend(built_tables)
        relationships.extend(built_relationships)

    if not tables:
        raise ValueError("All cleaned tables are empty; there is nothing to export to Power BI.")

    _add_cross_table_relationships(tables, relationships)
    measures = _build_measures(tables, relationships, results)
    dictionary = _build_data_dictionary(tables, relationships)
    kpis = _build_kpi_recommendations(tables, measures, relationships)
    visuals = _build_visual_recommendations(tables, measures, relationships)

    fact_tables = [table.name for table in tables if table.role == "Fact"]
    dimension_tables = [table.name for table in tables if table.role == "Dimension"]
    if not fact_tables:
        notes.append(
            "No numeric measures were found, so a star schema was not created. The cleaned table is exported as a flat table."
        )
    if dimension_tables:
        notes.append(
            "Dimensions use integer surrogate keys. Relationships are Many-to-One from the fact to the dimension."
        )
    notes.append(
        "Only measures with definitions that can be established from this dataset are included. "
        "Time-intelligence measures (YoY, YTD) are intentionally omitted until a date table is marked in Power BI."
    )
    notes.append(
        "CSV files in this package preserve raw values for Power BI import and are not neutralised for Excel formula execution. "
        "Import them into Power BI; do not double-click them in Excel."
    )

    model = PowerBIModel(
        tables=tables,
        relationships=relationships,
        measures=measures,
        data_dictionary=dictionary,
        kpi_recommendations=kpis,
        visual_recommendations=visuals,
        notes=notes,
        source_tables=source_tables,
        fact_tables=fact_tables,
        dimension_tables=dimension_tables,
    )
    return model


def _build_source_model(
    source_name: str,
    result: CleaningResult,
    *,
    multi_source: bool,
    max_dimensions: int,
    used_names: set[str],
) -> tuple[list[ModelTable], list[Relationship]]:
    cleaned = result.cleaned.copy()
    profiles = result.profile.column_profiles
    date_profiles = [profile for profile in profiles if profile.kind in {ColumnKind.DATE, ColumnKind.DATETIME}]
    measure_profiles = [
        profile for profile in profiles
        if profile.is_numeric_measure
        and profile.unique > 1
        and profile.missing_pct < 80
    ]
    dimension_profiles = [
        profile for profile in profiles
        if profile.role == "dimension"
        and profile.kind != ColumnKind.IDENTIFIER
        and 2 <= profile.unique <= 25
        and profile.missing_pct < 60
        and (not profile.sample or max(len(value) for value in profile.sample) <= 60)
    ]
    dimension_profiles = sorted(
        dimension_profiles,
        key=lambda profile: (profile.missing_pct, -profile.unique, profile.name),
    )[:max_dimensions]

    date_profile = date_profiles[0] if date_profiles else None
    can_star = bool(measure_profiles) and bool(dimension_profiles or date_profile) and len(cleaned) >= 1

    if not can_star:
        table_name = _unique_table_name(_beautify(source_name), used_names)
        table = ModelTable(
            name=table_name,
            dataframe=cleaned,
            role="Flat table",
            source_table=source_name,
            description="Cleaned source table exported without star-schema splitting.",
        )
        table.profiles = profile_dataframe(cleaned).column_profiles
        return [table], []

    tables: list[ModelTable] = []
    relationships: list[Relationship] = []
    fact = cleaned.copy()
    fact_name = _unique_table_name(f"Fact_{_beautify(source_name)}", used_names)
    tables_created: list[ModelTable] = []

    # Dimension tables from categorical columns.
    dimension_specs: list[tuple[ColumnProfile, str]] = []
    for profile in dimension_profiles:
        dim_name = _unique_table_name(
            f"Dim_{_beautify(profile.name)}" if not multi_source else f"Dim_{_beautify(source_name)}_{_beautify(profile.name)}",
            used_names,
        )
        dimension_specs.append((profile, dim_name))

    for profile, dim_name in dimension_specs:
        column = profile.name
        if column not in fact.columns:
            continue
        counts = fact[column].dropna().value_counts()
        if counts.empty:
            continue
        dim_frame = pd.DataFrame({column: counts.index.tolist()})
        key_column = _unique_column_name(f"{column}_key", [dim_frame.columns, fact.columns])
        dim_frame.insert(0, key_column, range(1, len(dim_frame) + 1))
        mapping = {value: index for index, value in enumerate(dim_frame[column].tolist(), start=1)}
        fact[key_column] = fact[column].map(mapping)
        fact = fact.drop(columns=[column])
        tables_created.append(
            ModelTable(
                name=dim_name,
                dataframe=dim_frame,
                role="Dimension",
                source_table=source_name,
                description=f"Distinct values of {column} from '{source_name}', with a surrogate key.",
            )
        )
        relationships.append(
            Relationship(
                from_table=fact_name,
                from_column=key_column,
                to_table=dim_name,
                to_column=key_column,
                cardinality="Many-to-One",
                cross_filter="Single",
                note=f"Fact rows link to one {column} value. Blank keys remain blank.",
            )
        )

    # Date dimension.
    if date_profile and date_profile.name in fact.columns:
        date_column = date_profile.name
        dates = pd.to_datetime(fact[date_column], errors="coerce").dropna()
        if not dates.empty:
            minimum = dates.min().normalize()
            maximum = dates.max().normalize()
            if (maximum - minimum).days <= 365 * 40:
                date_dim_name = _unique_table_name(
                    f"Dim_Date" if not multi_source else f"Dim_{_beautify(source_name)}_Date",
                    used_names,
                )
                calendar = pd.date_range(minimum, maximum, freq="D")
                date_frame = pd.DataFrame({"date": calendar})
                date_key = _unique_column_name("date_key", [fact.columns])
                date_frame.insert(0, date_key, date_frame["date"].dt.strftime("%Y%m%d").astype(int))
                date_frame["year"] = date_frame["date"].dt.year
                date_frame["quarter"] = "Q" + date_frame["date"].dt.quarter.astype(str)
                date_frame["month"] = date_frame["date"].dt.month
                date_frame["month_name"] = date_frame["date"].dt.strftime("%B")
                date_frame["day"] = date_frame["date"].dt.day
                date_frame["weekday"] = date_frame["date"].dt.strftime("%A")
                date_frame["year_month"] = date_frame["date"].dt.strftime("%Y-%m")

                fact_key = _unique_column_name(f"{date_column}_date_key", [fact.columns])
                fact[fact_key] = pd.to_datetime(fact[date_column], errors="coerce").dt.strftime("%Y%m%d")
                fact[fact_key] = pd.to_numeric(fact[fact_key], errors="coerce").astype("Int64")
                fact = fact.drop(columns=[date_column])
                tables_created.append(
                    ModelTable(
                        name=date_dim_name,
                        dataframe=date_frame,
                        role="Dimension",
                        source_table=source_name,
                        description=f"Daily calendar covering {minimum.date()} to {maximum.date()}.",
                    )
                )
                relationships.append(
                    Relationship(
                        from_table=fact_name,
                        from_column=fact_key,
                        to_table=date_dim_name,
                        to_column=date_key,
                        cardinality="Many-to-One",
                        cross_filter="Single",
                        note="Mark this table as a date table in Power BI before using time intelligence.",
                    )
                )
            else:
                # Extremely wide ranges would create an unjustifiably large
                # calendar; keep the original date column in the fact instead.
                pass

    # Add row identifier first for traceability.
    row_id = _unique_column_name("row_id", [fact.columns])
    fact.insert(0, row_id, range(1, len(fact) + 1))

    fact = fact.loc[:, ~fact.columns.duplicated()]
    tables_created.insert(
        0,
        ModelTable(
            name=fact_name,
            dataframe=fact,
            role="Fact",
            source_table=source_name,
            description="Measures, surrogate keys, and degenerate attributes at the row grain of the source.",
        ),
    )
    for table in tables_created:
        table.profiles = profile_dataframe(table.dataframe).column_profiles
    return tables_created, relationships


def _add_cross_table_relationships(tables: list[ModelTable], relationships: list[Relationship]) -> None:
    """Suggest relationships between flat tables that share a key-like column."""

    existing = {
        (relationship.from_table, relationship.from_column, relationship.to_table, relationship.to_column)
        for relationship in relationships
    }
    added = 0
    for index, left in enumerate(tables):
        for right in tables[index + 1 :]:
            if added >= 8:
                return
            shared = [column for column in left.dataframe.columns if column in right.dataframe.columns]
            for column in shared:
                if column.endswith("_key") or column == "row_id":
                    continue
                if not (
                    column.lower().endswith(("_id", "_code", "_no", "_number", "_key"))
                    or column.lower() in {"id", "code", "customer", "product", "region", "category"}
                ):
                    continue
                left_unique = left.dataframe[column].nunique(dropna=True) == len(left.dataframe)
                right_unique = right.dataframe[column].nunique(dropna=True) == len(right.dataframe)
                if not (left_unique or right_unique):
                    continue
                if right_unique:
                    candidate = (left.name, column, right.name, column)
                    cardinality = "Many-to-One"
                else:
                    candidate = (right.name, column, left.name, column)
                    cardinality = "Many-to-One"
                if candidate in existing:
                    continue
                relationships.append(
                    Relationship(
                        from_table=candidate[0],
                        from_column=candidate[1],
                        to_table=candidate[2],
                        to_column=candidate[3],
                        cardinality=cardinality,
                        cross_filter="Single",
                        note="Suggested from a shared column; verify uniqueness in Model view.",
                    )
                )
                existing.add(candidate)
                added += 1
                break


# ---------------------------------------------------------------------------
# DAX / KPI / data dictionary
# ---------------------------------------------------------------------------
def _build_measures(
    tables: list[ModelTable],
    relationships: list[Relationship],
    results: dict[str, CleaningResult],
) -> list[DaxMeasure]:
    measures: list[DaxMeasure] = []
    used_names: set[str] = set()
    measure_columns = 0
    max_measure_columns = 8

    for table in tables:
        if table.role == "Dimension":
            continue
        table_dax = _dax_table(table.name)
        measure_name = _unique_measure_name(f"{_humanize(table.name)} Row Count", used_names)
        measures.append(
            DaxMeasure(
                name=measure_name,
                table=table.name,
                expression=f"COUNTROWS({table_dax})",
                format_string="#,##0",
                description=f"Number of rows in {table.name}.",
            )
        )

        for profile in table.profiles:
            if measure_columns >= max_measure_columns:
                break
            if not profile.is_numeric_measure or profile.unique <= 1:
                continue
            column = profile.name
            if column.endswith("_key") or column == "row_id":
                continue
            column_dax = _dax_column(column)
            human = _humanize(column)
            mixed_currency = len(profile.currency_symbols) > 1
            symbol = profile.currency_symbols[0] if len(profile.currency_symbols) == 1 else None
            tokens = [token for token in column.lower().split("_") if token]
            non_additive_price = any(token in {"price", "rate", "unitprice"} for token in tokens)
            if profile.kind == ColumnKind.PERCENTAGE:
                format_string = '0.0"%"'
                expressions = [
                    ("Average", "AVERAGE", human),
                    ("Minimum", "MIN", human),
                    ("Maximum", "MAX", human),
                ]
            elif profile.kind == ColumnKind.CURRENCY and mixed_currency:
                format_string = "#,##0.00"
                expressions = [
                    ("Average", "AVERAGE", human),
                    ("Minimum", "MIN", human),
                    ("Maximum", "MAX", human),
                ]
            elif non_additive_price:
                format_string = f'"{symbol}"#,##0.00' if symbol else "#,##0.00"
                expressions = [
                    ("Average", "AVERAGE", human),
                    ("Minimum", "MIN", human),
                    ("Maximum", "MAX", human),
                ]
            else:
                format_string = f'"{symbol}"#,##0.00' if symbol else "#,##0.00"
                expressions = [
                    ("Total", "SUM", human),
                    ("Average", "AVERAGE", human),
                    ("Minimum", "MIN", human),
                    ("Maximum", "MAX", human),
                ]

            for prefix, function, human_name in expressions:
                dax_name = _unique_measure_name(f"{prefix} {human_name}", used_names)
                description = f"{prefix} of {column} in {table.name}."
                if profile.kind == ColumnKind.PERCENTAGE:
                    description += " Percentages are stored as percentage points."
                if mixed_currency:
                    description += " Mixed currency indicators detected; total is intentionally omitted."
                measures.append(
                    DaxMeasure(
                        name=dax_name,
                        table=table.name,
                        expression=f"{function}({table_dax}{column_dax})",
                        format_string=format_string,
                        description=description,
                    )
                )
            measure_columns += 1

    # Distinct counts on dimension attributes (not on surrogate keys).
    distinct_added = 0
    for table in tables:
        if distinct_added >= 6:
            break
        if table.role != "Dimension":
            continue
        attributes = [
            column for column in table.dataframe.columns
            if not column.endswith("_key") and column != "row_id"
        ]
        if not attributes:
            continue
        column = attributes[0]
        human = _humanize(column)
        dax_name = _unique_measure_name(f"Distinct {human}", used_names)
        measures.append(
            DaxMeasure(
                name=dax_name,
                table=table.name,
                expression=f"DISTINCTCOUNT({_dax_table(table.name)}{_dax_column(column)})",
                format_string="#,##0",
                description=f"Number of distinct {column} values in {table.name}.",
            )
        )
        distinct_added += 1

    return measures


def _build_kpi_recommendations(
    tables: list[ModelTable],
    measures: list[DaxMeasure],
    relationships: list[Relationship],
) -> pd.DataFrame:
    rows = []
    for measure in measures:
        if "Row Count" in measure.name:
            why = "Shows data volume and supports trend context."
            visual = "KPI card"
        elif measure.name.startswith("Total "):
            why = "Primary additive measure for totals and comparisons."
            visual = "KPI card + bar/line chart"
        elif measure.name.startswith("Average "):
            why = "Useful when totals are misleading, especially for percentages and mixed currencies."
            visual = "KPI card"
        elif measure.name.startswith("Distinct "):
            why = "Counts unique entities such as customers, products, or regions."
            visual = "KPI card"
        else:
            why = "Supports range and quality checks."
            visual = "Card or table"
        rows.append(
            {
                "KPI / measure": measure.name,
                "Table": measure.table,
                "DAX expression": measure.expression,
                "Why it matters": why,
                "Suggested visual": visual,
            }
        )
    return pd.DataFrame(rows)


def _build_visual_recommendations(
    tables: list[ModelTable],
    measures: list[DaxMeasure],
    relationships: list[Relationship],
) -> list[VisualRecommendation]:
    visuals: list[VisualRecommendation] = []
    total_measures = [measure for measure in measures if measure.name.startswith("Total ")]
    avg_measures = [measure for measure in measures if measure.name.startswith("Average ")]
    primary = (total_measures or avg_measures or measures)[0] if measures else None
    if primary:
        visuals.append(
            VisualRecommendation(
                visual="KPI cards",
                title="Headline metrics",
                fields=primary.name,
                why="Place the most important total or average at the top of the report page.",
            )
        )

    dimension_tables = [table for table in tables if table.role == "Dimension"]
    fact_tables = [table for table in tables if table.role == "Fact"]
    for dimension in dimension_tables[:3]:
        value_columns = [column for column in dimension.dataframe.columns if not column.endswith("_key")]
        if not value_columns:
            continue
        value = value_columns[0]
        measure = (total_measures or avg_measures or [None])[0]
        if measure:
            visuals.append(
                VisualRecommendation(
                    visual="Bar chart",
                    title=f"{measure.name} by {_humanize(value)}",
                    fields=f"{dimension.name}[{value}], {measure.name}",
                    why="Compares a measure across a small set of categories.",
                )
            )
        else:
            visuals.append(
                VisualRecommendation(
                    visual="Donut chart",
                    title=f"Record share by {_humanize(value)}",
                    fields=f"{dimension.name}[{value}]",
                    why="Shows composition for a low-cardinality category.",
                )
            )

    date_tables = [
        relation for relation in relationships
        if "date" in relation.to_table.lower()
    ]
    if date_tables and total_measures:
        date_table = next(table for table in tables if table.name == date_tables[0].to_table)
        if "date" in date_table.dataframe.columns:
            visuals.append(
                VisualRecommendation(
                    visual="Line chart",
                    title=f"{total_measures[0].name} over time",
                    fields=f"{date_table.name}[date], {total_measures[0].name}",
                    why="A daily date dimension supports a clean time axis and drill-down by year/month.",
                )
            )

    for dimension in dimension_tables[:3]:
        value_columns = [column for column in dimension.dataframe.columns if not column.endswith("_key")]
        if value_columns:
            visuals.append(
                VisualRecommendation(
                    visual="Slicer",
                    title=f"Filter by {_humanize(value_columns[0])}",
                    fields=f"{dimension.name}[{value_columns[0]}]",
                    why="Lets report viewers focus the page without duplicating visuals.",
                )
            )

    if fact_tables:
        visuals.append(
            VisualRecommendation(
                visual="Table / matrix",
                title="Detail table",
                fields="Fact rows with surrogate keys replaced by dimension attributes",
                why="Useful for reconciliation and drilled-through detail.",
            )
        )
    return visuals


def _build_data_dictionary(
    tables: list[ModelTable],
    relationships: list[Relationship],
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    fact_foreign_keys = {(relationship.from_table, relationship.from_column) for relationship in relationships}
    dim_primary_keys = {(relationship.to_table, relationship.to_column) for relationship in relationships}
    for table in tables:
        for profile in table.profiles:
            column = profile.name
            role = "Attribute"
            description = f"{profile.kind_label} column from '{table.source_table}'."
            if column == "row_id":
                role = "Row identifier"
                description = "Sequential identifier for traceability; not a business measure."
            elif (table.name, column) in dim_primary_keys or (
                table.role == "Dimension" and column.endswith("_key")
            ):
                role = "Primary key"
                description = f"Surrogate key for {table.name}; use this column on the one side of a relationship."
            elif (table.name, column) in fact_foreign_keys:
                role = "Foreign key"
                matching = next(
                    relationship for relationship in relationships
                    if relationship.from_table == table.name and relationship.from_column == column
                )
                description = f"Links to {matching.to_table}.{matching.to_column}."
            elif column.endswith("_key"):
                role = "Surrogate key"
                description = "Integer key used for relationships."
            elif profile.kind in {ColumnKind.DATE, ColumnKind.DATETIME}:
                role = "Date"
                description = f"Date attribute ({profile.kind_label}) from '{table.source_table}'."
            elif profile.is_numeric_measure and table.role != "Dimension":
                role = "Measure"
                if profile.kind == ColumnKind.CURRENCY:
                    description = f"Currency measure from '{table.source_table}'. Review currency context before summing."
                elif profile.kind == ColumnKind.PERCENTAGE:
                    description = f"Percentage measure from '{table.source_table}', stored as percentage points."
                else:
                    description = f"Numeric measure from '{table.source_table}'."
            elif table.role == "Dimension":
                role = "Dimension attribute"
                description = f"Descriptive attribute for grouping and filtering ({profile.kind_label})."

            source_column = column
            if column == "row_id":
                source_column = ""
            elif column.endswith("_key"):
                original = column[: -len("_key")]
                source_column = original if original in table.dataframe.columns else ""
            rows.append(
                {
                    "Table": table.name,
                    "Column": column,
                    "Role": role,
                    "Data type": profile.kind_label,
                    "Description": description,
                    "Source table": table.source_table,
                    "Source column": source_column,
                    "Missing": profile.missing,
                    "Unique": profile.unique,
                    "Examples": "; ".join(profile.sample[:3]),
                }
            )
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Export package
# ---------------------------------------------------------------------------
def build_powerbi_package(results: dict[str, CleaningResult]) -> PowerBIPackage:
    """Build the downloadable ZIP, workbook, and documentation."""

    model = prepare_model(results)
    workbook_bytes = _build_model_workbook(model)
    import_guide = _build_import_guide(model)
    dax_text = model.dax_text()
    dictionary_csv = dataframe_to_csv(model.data_dictionary, neutralise=False)
    relationships_csv = dataframe_to_csv(model.relationships_frame, neutralise=False)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("README_IMPORT.md", import_guide)
        archive.writestr("SECURITY_AND_LIMITATIONS.txt", _security_note())
        archive.writestr("CleanFlow_PowerBI_Model.xlsx", workbook_bytes)
        archive.writestr("dax_measures.dax", dax_text)
        archive.writestr("data_dictionary.csv", dictionary_csv)
        archive.writestr("relationships.csv", relationships_csv)
        archive.writestr("kpi_recommendations.csv", dataframe_to_csv(model.kpi_recommendations, neutralise=False))
        archive.writestr("visual_recommendations.csv", dataframe_to_csv(model.visuals_frame, neutralise=False))

        used_files: set[str] = set()
        file_map: dict[str, str] = {}
        for table in model.tables:
            filename = _unique_filename(_filename_safe(table.name), used_files)
            file_map[table.name] = f"tables/{filename}.csv"
            archive.writestr(f"tables/{filename}.csv", dataframe_to_csv(table.dataframe, neutralise=False))

        model_json = {
            "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "app_version": APP_VERSION,
            "source_tables": model.source_tables,
            "tables": [
                {
                    "name": table.name,
                    "role": table.role,
                    "rows": len(table.dataframe),
                    "columns": list(map(str, table.dataframe.columns)),
                    "source_table": table.source_table,
                    "description": table.description,
                    "csv": file_map.get(table.name, ""),
                }
                for table in model.tables
            ],
            "relationships": [
                {
                    "from_table": relationship.from_table,
                    "from_column": relationship.from_column,
                    "to_table": relationship.to_table,
                    "to_column": relationship.to_column,
                    "cardinality": relationship.cardinality,
                    "cross_filter": relationship.cross_filter,
                    "note": relationship.note,
                }
                for relationship in model.relationships
            ],
            "measures": [
                {
                    "name": measure.name,
                    "table": measure.table,
                    "expression": measure.expression,
                    "format_string": measure.format_string,
                    "description": measure.description,
                }
                for measure in model.measures
            ],
            "notes": model.notes,
        }
        archive.writestr("model.json", json.dumps(model_json, indent=2, default=str))

    return PowerBIPackage(
        model=model,
        zip_bytes=buffer.getvalue(),
        workbook_bytes=workbook_bytes,
        dax_text=dax_text,
        import_guide=import_guide,
        dictionary_csv=dictionary_csv,
        relationships_csv=relationships_csv,
    )


def _build_model_workbook(model: PowerBIModel) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        workbook = writer.book
        workbook.properties.creator = "CleanFlow AI"
        workbook.properties.title = "CleanFlow AI Power BI model"
        readme = _first_worksheet(workbook, "Read Me")
        _write_readme_sheet(readme, model)

        sheets_written = {"Read Me"}
        for frame, name in (
            (model.model_frame, "Model"),
            (model.relationships_frame, "Relationships"),
            (model.measures_frame, "DAX Measures"),
            (model.kpi_recommendations, "KPI Suggestions"),
            (model.visuals_frame, "Visual Ideas"),
            (model.data_dictionary, "Data Dictionary"),
        ):
            _write_frame(writer, frame, name, sheets_written)

        for table in model.tables:
            _write_frame(writer, table.dataframe, table.name, sheets_written, add_excel_table=True)
    return buffer.getvalue()


def _write_readme_sheet(sheet: Any, model: PowerBIModel) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 100
    sheet.merge_cells("A1:B1")
    cell = sheet["A1"]
    cell.value = "CleanFlow AI — Power BI Model Package"
    cell.font = Font(bold=True, size=20, color="FFFFFF")
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[1].height = 42

    sheet.merge_cells("A2:B2")
    cell = sheet["A2"]
    cell.value = "Power BI-ready data and a suggested model. This is not a .pbix report."
    cell.font = Font(bold=True, color=ROYAL, size=11)
    cell.alignment = Alignment(indent=1)

    row = 4
    for label, value in (
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
        ("Source tables", ", ".join(model.source_tables)),
        ("Fact tables", ", ".join(model.fact_tables) or "None"),
        ("Dimension tables", ", ".join(model.dimension_tables) or "None"),
        ("Model tables", str(len(model.tables))),
        ("Suggested measures", str(len(model.measures))),
        ("Relationships", str(len(model.relationships))),
        ("CleanFlow AI version", APP_VERSION),
    ):
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True, color=NAVY, size=10)
        sheet.cell(row=row, column=2, value=str(value)).font = BODY_FONT
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="IMPORTANT").font = Font(bold=True, color="B91C1C", size=11)
    row += 1
    for note in model.notes:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = sheet.cell(row=row, column=1, value="•  " + note)
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 30
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="NEXT STEPS").font = Font(bold=True, color=NAVY, size=11)
    row += 1
    for step in (
        "1. Open Power BI Desktop (free) and connect to the CSVs in tables/ or this workbook.",
        "2. Confirm column types, especially dates and decimal numbers.",
        "3. Create the relationships listed on the Relationships sheet.",
        "4. Mark Dim_Date as a date table if present.",
        "5. Add the measures from the DAX Measures sheet.",
        "6. Build visuals using the Visual Ideas sheet, then save your own .pbix file.",
    ):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell = sheet.cell(row=row, column=1, value=step)
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 22
        row += 1


def _write_frame(
    writer: pd.ExcelWriter,
    frame: pd.DataFrame,
    sheet_name: str,
    used: set[str],
    *,
    add_excel_table: bool = False,
) -> None:
    if frame is None or frame.shape[1] == 0:
        return
    safe = _unique_sheet_name(safe_sheet_name(sheet_name, "Sheet"), used)
    used.add(safe)
    prepared = frame.copy()
    for column in prepared.columns:
        if isinstance(prepared[column].dtype, pd.DatetimeTZDtype):
            prepared[column] = prepared[column].dt.tz_convert("UTC").dt.tz_localize(None)
        elif prepared[column].dtype == object:
            prepared[column] = prepared[column].map(_clean_cell)
    prepared.to_excel(writer, sheet_name=safe, index=False)
    sheet = writer.sheets[safe]
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 24
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    # Only touch cells that need formula protection; iterating millions of
    # cells is unnecessarily slow.
    risky_rows: dict[int, list[int]] = {}
    for column_index, column in enumerate(prepared.columns, start=1):
        mask = prepared[column].map(
            lambda value: isinstance(value, str) and is_formula_like(value)
        )
        for position in mask[mask].index.tolist():
            try:
                row_index = int(position) + 2
            except (TypeError, ValueError):
                continue
            risky_rows.setdefault(row_index, []).append(column_index)
    for row_index, columns in risky_rows.items():
        for column_index in columns:
            sheet.cell(row=row_index, column=column_index).data_type = "s"

    widths = []
    for column in prepared.columns:
        sample = prepared[column].head(100).astype(str).tolist()
        width = max([len(str(column)) + 2, 12] + [min(len(value), 50) for value in sample]) + 2
        widths.append(min(max(width, 12), 48))
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    reference = f"A1:{get_column_letter(prepared.shape[1])}{max(len(prepared) + 1, 2)}"
    if add_excel_table and len(prepared) > 0:
        table = Table(displayName=_unique_table_name_excel(safe, used), ref=reference)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        try:
            sheet.add_table(table)
        except Exception:
            sheet.auto_filter.ref = reference
    elif prepared.shape[1] > 0:
        sheet.auto_filter.ref = reference
    sheet.sheet_properties.tabColor = CYAN


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
    if isinstance(value, pd.Timestamp):
        if value.tzinfo is not None:
            value = value.tz_convert("UTC").tz_localize(None)
        return value.to_pydatetime()
    return value


# ---------------------------------------------------------------------------
# Documentation
# ---------------------------------------------------------------------------
def _first_worksheet(workbook: Workbook, title: str) -> Any:
    """Return the first worksheet without assuming ``workbook.active`` exists."""

    if workbook.sheetnames:
        sheet = workbook[workbook.sheetnames[0]]
        sheet.title = _unique_sheet_name(title, set(workbook.sheetnames[1:]))
    else:
        sheet = workbook.create_sheet(title)
    return sheet


def _build_import_guide(model: PowerBIModel) -> str:
    relationship_lines = "\n".join(
        f"| {relationship.from_table} | {relationship.from_column} | {relationship.to_table} | {relationship.to_column} | {relationship.cardinality} |"
        for relationship in model.relationships
    ) or "| — | — | — | — | No relationships suggested |"
    return f"""# CleanFlow AI — Power BI Import Guide

This package contains **Power BI-ready data and a suggested model**. It does **not** contain a `.pbix` report file. CleanFlow AI cannot create an arbitrary PBIX file with the free/open-source toolchain; you will build and save the report in Power BI Desktop.

## What is included

- `tables/*.csv` — one clean CSV per model table (UTF-8 with BOM).
- `CleanFlow_PowerBI_Model.xlsx` — all model tables plus documentation sheets.
- `model.json` — machine-readable table, relationship, and measure metadata.
- `relationships.csv`, `data_dictionary.csv`, `dax_measures.dax`.
- `kpi_recommendations.csv`, `visual_recommendations.csv`.
- `SECURITY_AND_LIMITATIONS.txt`.

## Suggested model

- Source tables: {", ".join(model.source_tables)}
- Fact tables: {", ".join(model.fact_tables) or "none"}
- Dimension tables: {", ".join(model.dimension_tables) or "none"}
- Measures suggested: {len(model.measures)}
- Relationships suggested: {len(model.relationships)}

| From table | From column | To table | To column | Cardinality |
|---|---|---|---|---|
{relationship_lines}

## Import steps (Windows)

1. Install **Power BI Desktop** from Microsoft (the free Windows desktop application). No paid license is required for local report authoring.
2. Extract this ZIP to a folder.
3. Open Power BI Desktop and choose **Get data → Text/CSV**.
4. Import each file in `tables/` one at a time. In the preview dialog choose:
   - **File origin:** 65001: Unicode (UTF-8)
   - **Delimiter:** Comma
   - **Data type detection:** Based on first 200 rows (or do not detect and set types manually)
5. Click **Transform Data** if you want to verify the column types in Power Query. Dates should be Date/Date/Time and numeric measures should be Decimal Number or Whole Number. Do not convert identifier columns to numbers; they are text to preserve leading zeros.
6. Rename each query to match its role in `model.json` if Power BI assigned a different name. This makes the suggested DAX easier to paste.
7. In **Model view**, create the relationships listed above:
   - Drag the fact key onto the matching dimension key.
   - Set cardinality to **Many to one (*:1)**.
   - Set cross-filter direction to **Single**.
   - Leave inactive relationships inactive until you have verified the data.
8. If `Dim_Date` exists, select it, then **Table tools → Mark as date table** and choose the `date` column. This enables time intelligence later.
9. In **Report view**, create measures using **Modeling → New measure** and paste each block from `dax_measures.dax`. Check the suggested format string on the measure's Format property.
10. Build visuals using `visual_recommendations.csv`. Save your report as a `.pbix` file on your own machine.

## Refresh workflow

1. Run the data through CleanFlow AI again when the source changes.
2. Download a new Power BI package.
3. Replace the CSV files (keep the same file names) or change the Power Query source path.
4. In Power BI Desktop choose **Refresh**. Relationships and measures are preserved in your `.pbix`.

## Important limitations

- Relationship suggestions are based on naming, uniqueness, and cardinality. They are **recommendations**, not a guaranteed correct model. Confirm them against your business rules.
- Currency symbols are recorded, but Power BI does not attach a currency to a number automatically. Set the format per measure; do not sum mixed currencies.
- Percentages are stored as percentage points (12.5%, not 0.125) and use the format string `0.0"%"`.
- A star schema is created only when at least one numeric measure and one useful dimension/date exist. Otherwise the cleaned table is exported flat.
- CleanFlow does not upload anything to Microsoft. All files are generated locally.
"""


def _security_note() -> str:
    return """CleanFlow AI — Power BI package security and limitations

- This package is generated locally from the cleaned dataset. No data is sent to
  Microsoft or any external AI service.
- CSV files preserve raw string values for Power BI import. They are not
  neutralised for spreadsheet formula execution. Import them into Power BI; do
  not open them directly in Excel if the source contains untrusted text.
- The Excel workbook forces cells that begin with formula characters to be
  stored as text.
- No .pbix file is generated. Microsoft does not provide a supported free API
  for arbitrary PBIX report generation. The package is a dataset and model
  blueprint that you open in Power BI Desktop.
- Model and DAX suggestions must be reviewed before production use.
- Uploaded data is processed in memory inside the CleanFlow app and is not
  written to a public web directory.
- Export files remain in your browser's downloads and are not retained by
  CleanFlow after the session ends. Cleaning history stores metadata only.
"""


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _dax_table(name: str) -> str:
    return "'" + str(name).replace("'", "''") + "'"


def _dax_column(name: str) -> str:
    return "[" + str(name).replace("]", "]]") + "]"


def _humanize(name: str) -> str:
    text = str(name).replace("_", " ").strip()
    return " ".join(word.capitalize() if word.islower() else word for word in text.split())


def _beautify(name: str) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", " ", str(name)).strip()
    if not text:
        return "Table"
    return "".join(word.capitalize() for word in text.split())


def _unique_table_name(base: str, used: set[str]) -> str:
    candidate = base[:60] or "Table"
    suffix = 2
    while candidate in used:
        candidate = f"{base[:52]} {suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _unique_column_name(base: str, existing: Iterable[Iterable[Any]]) -> str:
    columns: set[str] = set()
    for group in existing:
        columns.update(str(column) for column in group)
    candidate = base
    suffix = 2
    while candidate in columns:
        candidate = f"{base}_{suffix}"
        suffix += 1
    return candidate


def _unique_measure_name(base: str, used: set[str]) -> str:
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base} {suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _unique_sheet_name(name: str, used: set[str]) -> str:
    base = safe_sheet_name(name, "Sheet")
    candidate = base
    suffix = 2
    while candidate in used:
        tail = f" ({suffix})"
        candidate = base[: 31 - len(tail)] + tail
        suffix += 1
    return candidate


def _unique_table_name_excel(name: str, used: set[str]) -> str:
    cleaned = "".join(character for character in str(name) if character.isalnum() or character == "_")
    if not cleaned:
        cleaned = "Table"
    if not cleaned[0].isalpha() and cleaned[0] != "_":
        cleaned = "T_" + cleaned
    base = ("CleanFlow_" + cleaned)[:220]
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base[:210]}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _filename_safe(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", str(name)).strip("._-")
    return cleaned or "table"


def _unique_filename(base: str, used: set[str]) -> str:
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate
